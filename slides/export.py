#!/usr/bin/env python3
"""Generate HTML and an editable PPTX from content.json. Run render.py --pdf next."""
import html
import json
import re
import subprocess
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.util import Inches, Pt
from pptx.oxml.xmlchemy import OxmlElement
from style_deck import PALETTES, palette, title_box, summary_box

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
PART_NAMES = {1:'01  디자인 방향', 2:'02  여덟 가지 기법', 3:'03  Claude Code 실전'}
MEDIA = json.loads((HERE/'media/manifest.json').read_text())


def esc(v):
    return html.escape(str(v), quote=True)


def local(src):
    return '../'+src


def image_geometry(item):
    iw, ih = Image.open(ROOT/item['src']).size
    cx, cy, cw, ch = item.get('crop', [0, 0, iw, ih])
    assert 0 <= cx < iw and 0 <= cy < ih and cx+cw <= iw and cy+ch <= ih, item
    scale = min(item['w']/cw, item['h']/ch)
    w, h = cw*scale, ch*scale
    return (item['x']+(item['w']-w)/2, item['y']+(item['h']-h)/2, w, h, cx, cy, cw, ch, iw, ih, scale)


def html_image(item):
    x,y,w,h,cx,cy,cw,ch,iw,ih,s = image_geometry(item)
    source = item['src']
    media = MEDIA.get(source)
    poster = media['poster'] if media else source
    figure_match = re.search(r'figure-(\d+)', source)
    alt = item.get('label') or (f'Figure {int(figure_match[1])}' if figure_match else re.sub(r'[_-]', ' ', Path(source).stem))
    result = f'<figure class="figure-frame" style="left:{x:.2f}px;top:{y:.2f}px;width:{w:.2f}px;height:{h:.2f}px">'
    result += f'<img alt="{esc(alt)}" src="{local(poster)}" style="position:absolute;left:{-cx*s:.2f}px;top:{-cy*s:.2f}px;width:{iw*s:.2f}px;height:{ih*s:.2f}px" loading="eager">'
    if media:
        result += f'<video class="figure-video" data-src="{local(media["video"])}" muted loop playsinline preload="none" aria-label="움직이는 Figure" poster="{local(poster)}"></video>'
    result += '</figure>'
    if item.get('label'):
        result += f'<p class="figure-label" style="left:{item["x"]}px;top:{item["y"]+item["h"]+12}px;width:{item["w"]}px">{esc(item["label"])}</p>'
    return result


def html_text(t):
    style = f'left:{t["x"]}px;top:{t["y"]}px;width:{t["w"]}px;height:{t["h"]}px;font-size:{t.get("size",32)}px;color:{t.get("color","#151515")};font-weight:{t.get("weight",400)};text-align:{t.get("align","left")};line-height:{t.get("line_height",1.18)}'
    return f'<p class="diagram-text {esc(t.get("role",""))}" style="{style}">{esc(t["text"]).replace(chr(10),"<br>")}</p>'


def reading_elements(s):
    """Shared geometry keeps summary/prompt text editable and consistent in both formats."""
    shapes, texts = [], []
    colors = palette(s)
    if s.get('summary'):
        texts.append(dict(text=s['summary'],**summary_box(s),color=colors['muted'],role='summary'))
    if s.get('prompt'):
        if s.get('prompt_layout') == 'bottom':
            shapes.append(dict(kind='line',x=96,y=818,w=1728,h=0,stroke=colors['line']))
            texts.append(dict(text=s.get('prompt_label','프롬프트 요지'),x=96,y=848,w=180,h=92,size=20,weight=500,color=colors['faint'],role='prompt-label'))
            texts.append(dict(text=s['prompt'],x=302,y=840,w=1522,h=167,size=30,color=colors['ink'],line_height=1.35,role='prompt-body'))
        else:
            if s.get('prompt_panel') == 'ink':
                colors = PALETTES['ink']
                shapes.append(dict(kind='rect',x=1290,y=245,w=630,h=750,fill=colors['bg']))
            else:
                shapes.append(dict(kind='line',x=1285,y=284,w=0,h=640,stroke=colors['line']))
            texts.append(dict(text=s.get('prompt_label','프롬프트 요지'),x=1340,y=295,w=484,h=70,size=20,weight=500,color=colors['faint'],role='prompt-label'))
            texts.append(dict(text=s['prompt'],x=1340,y=363,w=484,h=590,size=30,color=colors['ink'],line_height=1.46,role='prompt-body'))
    return shapes, texts


def html_shape(t):
    if t['kind']=='line':
        return f'<svg class="diagram-line" style="left:{t["x"]}px;top:{t["y"]}px" width="{max(t["w"],2)}" height="{max(t["h"],2)}"><line x1="0" y1="0" x2="{t["w"]}" y2="{t["h"]}" stroke="{t.get("stroke","#aaa")}" stroke-width="2"/></svg>'
    return f'<div class="diagram-shape" style="left:{t["x"]}px;top:{t["y"]}px;width:{t["w"]}px;height:{t["h"]}px;background:{t.get("fill","transparent")};border:2px solid {t.get("stroke","transparent")};border-radius:{t.get("radius",0)}px"></div>'


def html_slide(s, n, total):
    part = s['part']
    title=title_box(s)
    title_style=f'left:{title["x"]}px;top:{title["y"]}px;width:{title["w"]}px;height:{title["h"]}px;font-size:{title["size"]}px;line-height:{title["line_height"]}'
    body = [f'<section class="slide minimal theme-{s.get("theme","paper")} layout-{s.get("layout","standard")}" data-hdr="none" data-part="{part}" aria-label="{n}. {esc(s["title"])}">']
    display_title = esc(s.get('display_title',s['title'])).replace('\n','<br>')
    body += [f'<p class="eyebrow">{esc(s.get("kicker", ""))}</p>',f'<h1 class="figure-title" style="{title_style}">{display_title}</h1>']
    reading_shapes, reading_texts = reading_elements(s)
    body += [html_shape(t) for t in s.get('shapes',[])+reading_shapes]
    body += [html_image(i) for i in s.get('images',[])]
    body += [html_text(t) for t in s.get('texts',[])+reading_texts]
    if s.get('caption'): body += [f'<p class="figure-caption">{esc(s["caption"])}</p>']
    body += [f'<footer class="figure-footer"><span>DESIGN EYE</span><span>{PART_NAMES[part]}</span><span>{n:02} / {total:02}</span></footer>']
    body += [f'<aside class="speaker-notes"><b>{esc(s["source"])}</b><br>{esc(s.get("notes","")).replace(chr(10),"<br>")}</aside>','</section>']
    return '\n'.join(body)


def rgb(hex_string):
    return RGBColor.from_string(hex_string.lstrip('#'))


def unit(px):
    return Inches(px/96)


def ppt_text(slide, value, x, y, w, h, size=32, color='#151515', weight=400, align='left', line_height=1.15):
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    shape = slide.shapes.add_textbox(unit(x),unit(y),unit(w),unit(h))
    tf=shape.text_frame
    tf.clear(); tf.word_wrap = True
    tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.TOP
    for k, line in enumerate(value.split('\n')):
        p=tf.paragraphs[0] if k==0 else tf.add_paragraph()
        # The family name resolves Korean reliably across PowerPoint and LibreOffice.
        family='Pretendard'
        p.text=line; p.font.name=family; p.font.size=Pt(size*.75)
        p.font.bold=weight>=600; p.font.color.rgb=rgb(color)
        p.space_before=Pt(0); p.space_after=Pt(0)
        p.alignment={'left':PP_ALIGN.LEFT,'center':PP_ALIGN.CENTER,'right':PP_ALIGN.RIGHT}.get(align,PP_ALIGN.LEFT)
        p.line_spacing=line_height
        for run in p.runs:
            props=run._r.get_or_add_rPr(); props.set('lang','ko-KR')
            ea=OxmlElement('a:ea'); ea.set('typeface',family); props.append(ea)
    return shape


def ppt_image(slide, item, colors):
    x,y,w,h,cx,cy,cw,ch,iw,ih,scale=image_geometry(item)
    src = MEDIA.get(item['src'],{}).get('poster',item['src'])
    pic=slide.shapes.add_picture(str(ROOT/src),unit(x),unit(y),width=unit(w),height=unit(h))
    pic.crop_left=cx/iw; pic.crop_top=cy/ih
    pic.crop_right=(iw-cx-cw)/iw; pic.crop_bottom=(ih-cy-ch)/ih
    pic.name=item.get('label') or Path(item['src']).stem
    if item.get('label'):
        ppt_text(slide,item['label'],item['x'],item['y']+item['h']+12,item['w'],42,22,colors['muted'],500,'center')


def ppt_shape(slide,t):
    if t['kind']=='line':
        s=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,unit(t['x']),unit(t['y']),unit(t['x']+t['w']),unit(t['y']+t['h']))
    else:
        kind=MSO_SHAPE.ROUNDED_RECTANGLE if t.get('radius') else MSO_SHAPE.RECTANGLE
        s=slide.shapes.add_shape(kind,unit(t['x']),unit(t['y']),unit(t['w']),unit(t['h']))
        fill=t.get('fill','transparent')
        if fill=='transparent': s.fill.background()
        else: s.fill.solid(); s.fill.fore_color.rgb=rgb(fill)
    # Clear theme effects for both shapes and rules.
    s._element.spPr.append(OxmlElement('a:effectLst'))
    from pptx.oxml.ns import qn
    effect_ref = s._element.find('.//'+qn('a:effectRef'))
    if effect_ref is not None: effect_ref.set('idx','0')
    stroke=t.get('stroke','transparent')
    if stroke=='transparent': s.line.fill.background()
    else: s.line.color.rgb=rgb(stroke); s.line.width=Pt(1.5)


def export_pptx(slides):
    pres=Presentation(); pres.slide_width=unit(1920); pres.slide_height=unit(1080)
    pres.core_properties.title='Claude의 디자인 안목'
    pres.core_properties.subject='SLIDE1 → NOTE2 → HOW_DESIGN_FROM_CC'
    for n,s in enumerate(slides,1):
        p=pres.slides.add_slide(pres.slide_layouts[6])
        colors=palette(s)
        p.background.fill.solid(); p.background.fill.fore_color.rgb=rgb(colors['bg'])
        kicker_y=664 if s.get('layout')=='case-opener' else 40
        ppt_text(p,s.get('kicker',''),96,kicker_y,1700,30,18,colors['faint'],500)
        ppt_text(p,s.get('display_title',s['title']),**title_box(s),color=colors['ink'])
        reading_shapes, reading_texts = reading_elements(s)
        for shape in s.get('shapes',[])+reading_shapes: ppt_shape(p,shape)
        for item in s.get('images',[]): ppt_image(p,item,colors)
        for t in s.get('texts',[])+reading_texts: ppt_text(p,t['text'],t['x'],t['y'],t['w'],t['h'],t.get('size',32),t.get('color','#151515'),t.get('weight',400),t.get('align','left'),t.get('line_height',1.18))
        if s.get('caption'): ppt_text(p,s['caption'],96,974,1728,45,27,'#666661',400,'center')
        ppt_text(p,'DESIGN EYE',96,1030,500,30,17,colors['faint'],500)
        ppt_text(p,PART_NAMES[s['part']],580,1030,760,30,17,colors['faint'],400,'center')
        ppt_text(p,f'{n:02} / {len(slides):02}',1630,1030,194,30,17,colors['faint'],400,'right')
        p.notes_slide.notes_text_frame.text=s['source']+'\n\n'+s.get('notes','')+'\n\n이미지 출처:\n'+'\n'.join(i['src'] for i in s.get('images',[]))
    path=HERE/'design-eye.pptx'; pres.save(path)
    print(f'{path.name}: {len(slides)} slides, editable images and text')


def main():
    subprocess.run([sys.executable,str(HERE/'compose.py')],check=True)
    content=json.loads((HERE/'content.json').read_text()); slides=content['slides']
    parts={1:[],2:[],3:[]}
    for n,s in enumerate(slides,1): parts[s['part']].append(html_slide(s,n,len(slides)))
    for part, body in parts.items():
        (HERE/'parts'/f'{part:02}-figures.html').write_text('\n\n'.join(body)+'\n')
    subprocess.run([sys.executable,str(HERE/'build.py')],check=True)
    export_pptx(slides)


if __name__=='__main__': main()
